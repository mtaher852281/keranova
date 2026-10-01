"""Build and validate seven editable KeraNova protocol documents.

Inputs are versioned JSON drafts in protocol_sources. Outputs default to the
ignored generated directory, with a manifest for verification before promotion.
This builder does not read patient records, train a model, or calculate results.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCES = ROOT / "protocol_sources"
OUTPUT = ROOT / "generated"
FILENAMES = {
    1: "1. Methodology.docx",
    2: "2. Data_Collection_Sheet.docx",
    3: "3. Data_Variable_Mapping_Sheet.docx",
    4: "4. Data_Cleaning_Rules.docx",
    5: "5. Data_Transformation_Template.docx",
    6: "6. Statistical_Plan.docx",
    7: "7. Dummy_Tables_and_Figures.docx",
}


def load_drafts() -> dict[int, dict[str, Any]]:
    """Read both reviewed content files into a document-number map.

    Returns:
        Draft content for each of the seven requested documents.
    Raises:
        ValueError: If the input document mapping is incomplete or ambiguous.
        OSError: If an input file cannot be read.
    """
    drafts: dict[int, dict[str, Any]] = {}
    for name in ("methods_statistics.json", "data_workflow.json"):
        content = json.loads((SOURCES / name).read_text(encoding="utf-8"))
        entries = content.get("documents", content)
        if isinstance(entries, list):
            entries = {str(item.get("number", item.get("id"))): item for item in entries}
        for key, item in entries.items():
            if not isinstance(item, dict) or "title" not in item:
                continue
            number = int(str(key).split(".")[0])
            if number in drafts:
                raise ValueError(f"Duplicate document number {number}")
            drafts[number] = item
    if set(drafts) != set(FILENAMES):
        raise ValueError(f"Expected seven drafts, received {sorted(drafts)}")
    return drafts


def configure_document(doc: Any) -> None:
    """Apply readable page, paragraph and pagination defaults.

    Args:
        doc: Editable python-docx document.
    Returns:
        None. Mutates the in-memory document only.
    """
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Title", 21), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[style_name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    doc.styles["Title"].paragraph_format.space_before = Pt(0)
    for border in list(doc.styles.element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("KeraNova  |  ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.author = "KeraNova Research Team"
    doc.core_properties.subject = "Keratoconus cohort and explainable outcome prediction"
    doc.core_properties.keywords = "KeraNova, keratoconus, research protocol"


def style_cell(cell: Any, text: Any, header: bool = False) -> None:
    """Set a table cell's text, padding and legible typography.

    Args:
        cell: Target table cell.
        text: Display text, converted to a string.
        header: Whether to use header emphasis.
    Returns:
        None. Mutates the target cell.
    """
    cell.text = str(text if text is not None else "")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name, run.font.size = "Calibri", Pt(9)
            run.font.bold = header
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), "70")
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    props.append(margins)
    if header:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8F1F3")
        props.append(shading)


def add_table(doc: Any, spec: dict[str, Any]) -> None:
    """Insert a table with repeating headers and rows that do not split.

    Args:
        doc: Editable document.
        spec: Table specification containing headers and rows.
    Returns:
        None. Adds content to the document.
    Raises:
        ValueError: If a row width differs from its header width.
    """
    headers = spec.get("headers", spec.get("columns", []))
    if not headers:
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style, table.alignment = "Table Grid", WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = spec.get("widths", [6.97 / len(headers)] * len(headers))
    if headers[0] == "CSV column":
        widths = [0.65, 1.95, 1.65, 2.72]
    elif len(headers) == 3:
        widths = [1.5, 2.0, 3.47]
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(float(width))
    for index, (cell, label) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(float(widths[index]))
        style_cell(cell, label, True)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in spec.get("rows", []):
        if len(values) != len(headers):
            raise ValueError(f"Table row has {len(values)} cells, expected {len(headers)}")
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(float(widths[index]))
            style_cell(cell, value)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_section(doc: Any, spec: dict[str, Any]) -> None:
    """Append one content section with paragraphs, lists and tables.

    Args:
        doc: Editable document.
        spec: Structured section content.
    Returns:
        None. Adds content to the document.
    """
    if spec.get("page_break"):
        doc.add_page_break()
    if spec.get("heading"):
        heading = re.sub(r"^(\d+)\.\s*", r"\1 ", spec["heading"])
        heading = re.sub(r"^(Table\s+\d+)\.\s*", r"\1 ", heading)
        doc.add_heading(heading, level=spec.get("level", 1))
    for text in spec.get("paragraphs", []):
        doc.add_paragraph(str(text))
    for text in spec.get("bullets", []):
        doc.add_paragraph(str(text), style="List Bullet")
    if spec.get("table"):
        add_table(doc, spec["table"])
    for table in spec.get("tables", []):
        add_table(doc, table)
    for text in spec.get("notes", []):
        paragraph = doc.add_paragraph(str(text))
        for run in paragraph.runs:
            run.font.size = Pt(9)


def all_text(doc: Any) -> str:
    """Return body paragraph and table text for content validation.

    Args:
        doc: Loaded Word document.
    Returns:
        All body text joined by newlines.
    """
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def validate_document(path: Path) -> dict[str, Any]:
    """Reopen a DOCX and verify basic research-content and OOXML integrity.

    Args:
        path: Saved document path.
    Returns:
        File manifest record with counts and SHA256.
    Raises:
        ValueError: If obsolete glaucoma terms or empty content are found.
    """
    doc = Document(path)
    text = all_text(doc)
    forbidden = ("juvenile open-angle", "deep sclerectomy", "kaplan-meier", "5-21 mmhg")
    found = [term for term in forbidden if term in text.lower()]
    if found or len(text) < 500:
        raise ValueError(f"Content validation failed for {path.name}: {found}")
    if ":codex-" in text or "turn12search" in text:
        raise ValueError(f"Tool tokens leaked into {path.name}")
    return {"file": path.name, "characters": len(text), "tables": len(doc.tables),
            "paragraphs": len(doc.paragraphs), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def main() -> None:
    """Build all reviewed drafts and save their verification manifest.

    Returns:
    None. Writes seven DOCX files and a JSON manifest into the chosen output.
    Raises:
        OSError or ValueError: If drafts are invalid or outputs cannot be saved.
    """
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir", type=Path, default=OUTPUT,
        help="Output directory; defaults to generated beside the README.",
    )
    output_dir = parser.parse_args().output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest = []
    for number, content in sorted(load_drafts().items()):
        doc = Document()
        configure_document(doc)
        title = re.sub(r"[^\w\s]", " ", content["title"])
        doc.add_paragraph(re.sub(r"\s+", " ", title).strip(), style="Title")
        if content.get("subtitle"):
            subtitle = content["subtitle"].split("|")[0].strip()
            doc.add_paragraph(subtitle)
        doc.add_paragraph("KeraNova Research Team  |  Version 1.0  |  1 October 2026")
        for index, section in enumerate(content["sections"]):
            # Each mapping domain and result shell starts on its own page.
            if number == 3:
                section["page_break"] = 2 <= index <= 7
            if number == 7:
                section["page_break"] = 1 <= index <= 9
            add_section(doc, section)
        path = output_dir / FILENAMES[number]
        doc.save(path)
        manifest.append(validate_document(path))
    (output_dir / "document_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
