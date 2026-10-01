"""Audit KeraNova source data and save evidence without modifying source files."""

import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"C:\Users\user\Desktop\Hackathon")
MISSING = {"", "-", ".", "na", "n/a", "nd", "nan"}


def populated(value: str) -> bool:
    """Return whether a source string is populated, preserving explicit negatives."""
    return value.strip().casefold() not in MISSING


def numeric(value: str, field: str = "") -> float | None:
    """Parse a numeric source value for auditing only.

    Args:
        value: Original source text with optional units or trailing positive sign.
        field: Context permitting plano only in a sphere field.
    Returns:
        Parsed number or None for missing, nonnumeric or ambiguous input.
    """
    text = value.strip()
    if not populated(text):
        return None
    if field == "sphere" and text.casefold() in {"pl", "plano"}:
        return 0.0
    text = re.sub(r"\s*(?:diopters|D|μm|µm|um|y|years)\s*$", "", text, flags=re.I)
    text = text.replace(" ", "").replace(",", ".")
    if re.fullmatch(r"\d+(?:\.\d+)?\+", text):
        text = "+" + text[:-1]
    return float(text) if re.fullmatch(r"[+-]?\d+(?:\.\d+)?", text) else None


def audit_csv() -> dict:
    """Read CSV, calculate source counts and review candidates, return audit evidence."""
    source = ROOT / "dataset.csv"
    with source.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    headers = rows[0]
    section_rows = [(i, row[0]) for i, row in enumerate(rows[1:], 2) if row[0].strip() and not any(populated(v) for v in row[1:])]
    data = [(i, row) for i, row in enumerate(rows[1:], 2) if i not in {item[0] for item in section_rows}]
    nonmissing = [sum(populated(row[j]) for _, row in data) for j in range(len(headers))]
    corrected = {headers[j]: [{"source_row": i, "raw": row[j]} for i, row in data if populated(row[j])] for j in (9, 10, 11, 12)}
    mrse = []
    for start, interval in ((14, "baseline"), (17, "1_month_label"), (20, "6_month_label"), (23, "12_month_label")):
        for i, row in data:
            values = [numeric(row[start], "sphere"), numeric(row[start + 1]), numeric(row[start + 2])]
            if None in values:
                continue
            expected = values[0] + values[1] / 2
            if abs(values[2] - expected) > 0.0200001:
                mrse.append({"source_row": i, "interval": interval, "sphere_raw": row[start], "cylinder_raw": row[start+1], "mrse_raw": row[start+2], "mrse_calculated": expected, "difference": values[2] - expected})
    return {
        "source": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "csv_data_rows_before_classification": len(rows)-1, "eye_record_rows": len(data),
        "column_count": len(headers), "section_source_rows": [i for i,_ in section_rows],
        "headers": [{"column": j+1, "exact_header": h, "escaped_header": json.dumps(h), "nonmissing_eye_records": nonmissing[j]} for j,h in enumerate(headers)],
        "treatment_counts_raw": dict(Counter(row[38].strip() for _,row in data)),
        "sex_counts_raw": dict(Counter(row[1] for _,row in data)),
        "age_under18": [{"source_row": i,"raw": row[2]} for i,row in data if numeric(row[2]) is not None and numeric(row[2])<18],
        "corrected_va_entries": corrected,
        "baseline_and_12mo_cdva_populated_rows": [i for i,row in data if populated(row[9]) and populated(row[12])],
        "primary_candidate_fields_complete_raw_rows": [i for i,row in data if all(populated(row[j]) for j in (1,2,9,12,29,30,38))],
        "mrse_inconsistency_review_candidates": mrse,
        "haze_counts_raw": dict(Counter(row[41].strip() for _,row in data)),
        "exact_duplicate_source_row_pairs": [[data[a][0],data[b][0]] for a in range(len(data)) for b in range(a+1,len(data)) if data[a][1]==data[b][1]],
        "limitations": ["Eye records are source rows and do not establish unique patients or eligible eyes.","Source labels do not verify topography guided PRK or accelerated CXL.","No patient codes, index treatment dates, visit dates, imaging references or confirmed scan quality.","Postoperative topography fields have no named interval and must not be mapped to 12 months.","Nonmissing counts include qualified visual acuities and other values requiring source review.","Age timing is unspecified; entries under 18 require verification of age at index treatment."]
    }


def audit_workbook() -> dict:
    """Read the source workbook and return sheet dimensions and populated row evidence."""
    source = ROOT / "others" / "25240-R Database.xlsx"
    workbook = load_workbook(source, data_only=False, read_only=True)
    sheets = []
    for sheet in workbook:
        rows = list(sheet.iter_rows(values_only=True))
        nonempty = [(i, row) for i,row in enumerate(rows,1) if any(v is not None for v in row)]
        sheets.append({"sheet": sheet.title,"max_row": sheet.max_row,"max_column": sheet.max_column,"nonempty_row_count": len(nonempty),"first_rows": [[i,list(row)] for i,row in nonempty[:3]]})
    workbook.close()
    return {"source": str(source.relative_to(ROOT)),"sheets": sheets}


if __name__ == "__main__":
    result = {"audit_date": "2026-10-01", "csv": audit_csv(), "workbook": audit_workbook()}
    output = ROOT / "_keranova_build" / "data_audit.json"
    output.write_text(json.dumps(result,indent=2,ensure_ascii=False,default=str),encoding="utf-8")
    print(json.dumps(result,ensure_ascii=True,default=str))
